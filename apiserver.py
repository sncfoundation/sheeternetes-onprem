#!/usr/bin/env python3
"""
Sheeternetes on-prem apiserver — a LOCAL, offline control plane over a desktop
spreadsheet (Excel .xlsx). Same verb contract as the Google Apps Script apiserver,
so kubelet.sh / skctl point at it unchanged. Air-gap-friendly: no internet required.

  WORKBOOK=cluster.xlsx TOKEN=secret python3 apiserver.py            # serve on :8787
  GET  /?token=..&kind=pods|nodes|deployments|events   -> {"items":[...]}
  POST /  {"token":..,"action":"apply|scale|delete|cordon|uncordon|drain|migrate|label|taint", ...}
  POST /  {"token":..,"node":..,"ip":..,"cpu_total":..,"mem_total":..,"pods":[...]}  # kubelet heartbeat

The apiserver also SCHEDULES: on each heartbeat it bin-packs Deployment replicas onto
Ready, schedulable nodes by cpu_req/mem_req (a pod that fits nowhere is Unschedulable),
keeps pods sticky to their node, ages out silent nodes (failover), and honors
cordon/drain. The scheduler core is the pure function `schedule()` below.

Auth: every request carries a shared token. Additionally, if SIGNING_KEY is set, POSTs
must be HMAC-SHA256 signed (X-SNCF-Timestamp + X-SNCF-Signature) within SIGN_TTL seconds —
tamper- and replay-resistant. See bridge.py, which signs its cross-substrate payloads.

Requires: openpyxl  (pip install openpyxl). The spreadsheet is the store; this process
is the apiserver. See bridge.py for hybrid federation with Google Sheets.
"""
import hashlib, hmac, json, os, time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs

WORKBOOK = os.environ.get("WORKBOOK", "cluster.xlsx")
TOKEN = os.environ.get("TOKEN", "CHANGE_ME_super_secret")
PORT = int(os.environ.get("PORT", "8787"))
NODE_TTL = int(os.environ.get("NODE_TTL", "30"))   # seconds before a silent node is NotReady
SIGNING_KEY = os.environ.get("SIGNING_KEY", "")    # if set, POSTs must carry a valid HMAC
SIGN_TTL = int(os.environ.get("SIGN_TTL", "300"))  # max clock skew (s) for a signed request
TABS = {
    "Deployments": ["name", "image", "replicas", "cpu_req", "mem_req", "command", "node_selector", "tolerations"],
    "Nodes": ["name", "ip", "cpu_total", "cpu_used", "mem_total", "status", "last_heartbeat", "schedulable", "labels", "taints"],
    "Pods": ["name", "deployment", "node", "phase", "container_id"],
    "Events": ["ts", "kind", "object", "message"],
}

# ---------------------------------------------------------------- workbook I/O

def _wb():
    import openpyxl
    if os.path.exists(WORKBOOK):
        wb = openpyxl.load_workbook(WORKBOOK)
        _ensure_schema(wb)
        return wb
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for tab, headers in TABS.items():
        ws = wb.create_sheet(tab); ws.append(headers)
    wb.save(WORKBOOK); return wb

def _ensure_schema(wb):
    """Upgrade older workbooks: create missing tabs and append missing trailing
    columns (e.g. the newer Nodes.schedulable), backfilling a sane default."""
    changed = False
    defaults = {"schedulable": True}
    for tab, headers in TABS.items():
        if tab not in wb.sheetnames:
            wb.create_sheet(tab).append(headers); changed = True; continue
        ws = wb[tab]
        have = [str(c.value) for c in ws[1]] if ws.max_row else []
        for h in headers:
            if h not in have:
                col = len(have) + 1
                ws.cell(row=1, column=col, value=h)
                for r in range(2, ws.max_row + 1):
                    if ws.cell(row=r, column=1).value not in (None, ""):
                        ws.cell(row=r, column=col, value=defaults.get(h, ""))
                have.append(h); changed = True
    if changed: wb.save(WORKBOOK)

def _dicts(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return []
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if r and r[0] not in (None, "")]

def read_tab(kind):
    tab = {"pods": "Pods", "nodes": "Nodes", "deployments": "Deployments", "events": "Events"}.get(kind)
    if not tab: return None
    return _dicts(_wb()[tab])

def _int(v, default=0):
    try: return int(float(v))
    except (TypeError, ValueError): return default

def _truthy(v, default=True):
    if v in (None, ""): return default
    return str(v).strip().upper() not in ("FALSE", "0", "NO")

def sign(key, ts, body):
    """HMAC-SHA256 over 'timestamp.body' — the wire signature. Pure, so the bridge
    can produce byte-identical signatures and both sides can unit-test it."""
    if isinstance(body, str): body = body.encode()
    msg = str(ts).encode() + b"." + body
    return hmac.new(key.encode(), msg, hashlib.sha256).hexdigest()

def verify(key, ts, sig, body, now, ttl):
    """Constant-time verify with a bounded clock skew (anti-replay)."""
    if not (key and ts and sig): return False
    try:
        if abs(now - int(ts)) > ttl: return False
    except (TypeError, ValueError):
        return False
    return hmac.compare_digest(sign(key, ts, body), sig)

def parse_kv(s):
    """'disk=ssd,zone=a' -> {'disk':'ssd','zone':'a'}  (labels / node_selector)."""
    out = {}
    for part in str(s or "").split(","):
        part = part.strip()
        if "=" in part:
            k, v = part.split("=", 1); out[k.strip()] = v.strip()
    return out

def parse_taints(s):
    """'gpu=true:NoSchedule,zone=a' -> [('gpu','true','NoSchedule'),('zone','a','NoSchedule')]."""
    out = []
    for part in str(s or "").split(","):
        part = part.strip()
        if not part: continue
        kv, _, eff = part.partition(":")
        eff = eff.strip() or "NoSchedule"
        k, _, v = kv.partition("="); out.append((k.strip(), v.strip(), eff))
    return out

def parse_tolerations(s):
    """'gpu=true,zone' -> {'gpu=true','zone'}  (matches a taint by 'k=v' or bare 'k')."""
    return {p.strip() for p in str(s or "").split(",") if p.strip()}

def _tolerates(tols, taints):
    for k, v, eff in taints:
        if eff == "NoSchedule" and f"{k}={v}" not in tols and k not in tols:
            return False
    return True

def _matches(node, dep):
    """A pod's node_selector must be a subset of the node's labels, and the pod
    must tolerate every NoSchedule taint on the node."""
    labels = node.get("labels") or {}
    for k, v in (dep.get("node_selector") or {}).items():
        if labels.get(k) != v:
            return False
    return _tolerates(dep.get("tolerations") or set(), node.get("taints") or [])

# ------------------------------------------------------------- CRUD (skctl verbs)

def upsert_deployment(dep):
    wb = _wb(); ws = wb["Deployments"]; headers = TABS["Deployments"]
    name = dep.get("name")
    for row in ws.iter_rows(min_row=2):
        if row[0].value == name:
            for i, h in enumerate(headers): row[i].value = dep.get(h, row[i].value)
            wb.save(WORKBOOK); return "updated"
    ws.append([dep.get(h, "") for h in headers]); wb.save(WORKBOOK); return "created"

def scale(name, replicas):
    wb = _wb(); ws = wb["Deployments"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == name: row[2].value = replicas; wb.save(WORKBOOK); return "scaled"
    return "not found"

def delete(name):
    wb = _wb(); ws = wb["Deployments"]
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if row[0].value == name: ws.delete_rows(i, 1); wb.save(WORKBOOK); return "deleted"
    return "not found"

def set_schedulable(node, value):
    wb = _wb(); ns = wb["Nodes"]; col = TABS["Nodes"].index("schedulable")
    for row in ns.iter_rows(min_row=2):
        if row[0].value == node:
            row[col].value = bool(value)
            _reschedule(wb); wb.save(WORKBOOK)
            return "uncordoned" if value else "cordoned"
    return "not found"

def drain(node):
    """Cordon the node and move its pods off now (they reschedule onto survivors)."""
    wb = _wb(); ns = wb["Nodes"]; col = TABS["Nodes"].index("schedulable")
    found = False
    for row in ns.iter_rows(min_row=2):
        if row[0].value == node: row[col].value = False; found = True
    if not found: return "not found"
    _reschedule(wb, exclude={node}); wb.save(WORKBOOK)
    return "drained"

def migrate(pod, target):
    """Pin a pod to a target node (make-before-break: the target kubelet starts it,
    the source kubelet then sees it as Terminating and stops it)."""
    wb = _wb(); now = int(time.time())
    nodes = _load_nodes(wb["Nodes"], now)
    tgt = next((n for n in nodes if n["name"] == target), None)
    if not tgt: return "unknown node"
    if not (tgt["fresh"] and tgt["schedulable"]): return "target not schedulable"
    ps = wb["Pods"]
    for row in ps.iter_rows(min_row=2):
        if row[0].value == pod:
            row[2].value = target; wb.save(WORKBOOK); return f"migrating {pod} -> {target}"
    return "pod not found"

def _edit_node_csv(node, column, mutate):
    wb = _wb(); ns = wb["Nodes"]; col = TABS["Nodes"].index(column)
    for row in ns.iter_rows(min_row=2):
        if row[0].value == node:
            row[col].value = mutate(str(row[col].value or ""))
            _reschedule(wb); wb.save(WORKBOOK); return "ok"
    return "not found"

def label_node(node, spec):
    """label_node('a','disk=ssd')  sets a label; 'disk-' removes it."""
    def mut(cur):
        kv = parse_kv(cur)
        if spec.endswith("-"):
            kv.pop(spec[:-1].strip(), None)
        else:
            k, _, v = spec.partition("="); kv[k.strip()] = v.strip()
        return ",".join(f"{k}={v}" for k, v in kv.items())
    return _edit_node_csv(node, "labels", mut)

def taint_node(node, spec):
    """taint_node('a','gpu=true:NoSchedule')  adds a taint; 'gpu-' removes by key."""
    def mut(cur):
        taints = parse_taints(cur)
        if spec.endswith("-"):
            key = spec[:-1].strip()
            taints = [t for t in taints if t[0] != key]
        else:
            for k, v, eff in parse_taints(spec):
                taints = [t for t in taints if t[0] != k] + [(k, v, eff)]
        return ",".join(f"{k}={v}:{eff}" for k, v, eff in taints)
    return _edit_node_csv(node, "taints", mut)

# ------------------------------------------------------------------ scheduler

def schedule(deployments, nodes, existing, exclude=frozenset()):
    """Pure scheduler — no I/O, fully unit-testable.

    deployments: [{name,image,replicas,cpu_req,mem_req,command,node_selector,tolerations}]
    nodes:       [{name,cpu_total,mem_total,fresh(bool),schedulable(bool),labels,taints}]
    existing:    {podname: {"node": ...}}   (for sticky placement)
    exclude:     node names to evict from (drain) — their pods are re-placed.

    Affinity: a pod is only placed on a node whose labels are a superset of the
    deployment's node_selector and whose NoSchedule taints the pod tolerates.

    Returns (desired, alloc):
      desired = {podname: {deployment,node,image,command,cpu_req,mem_req}}
                node == "" means Unschedulable (fits nowhere / no capacity).
      alloc   = {nodename: {"cpu": millicores, "mem": MiB}}  placed load per node.
    Placement is resource-aware best-effort spread: a pod stays on its current node
    if that node is fresh, not excluded, and still has room; otherwise it lands on the
    fresh+schedulable node with the most free CPU that can fit it.
    """
    keepable = {n["name"] for n in nodes if n["fresh"]} - set(exclude)
    by_name = {n["name"]: n for n in nodes}
    schedulable_fresh = [n for n in nodes if n["fresh"] and n["schedulable"] and n["name"] not in exclude]
    cap = {n["name"]: (n["cpu_total"], n["mem_total"]) for n in nodes}
    alloc = {n["name"]: {"cpu": 0, "mem": 0} for n in nodes}

    def fits(name, cpu, mem):
        ct, mt = cap.get(name, (0, 0))
        return alloc[name]["cpu"] + cpu <= ct and alloc[name]["mem"] + mem <= mt

    def place(name, cpu, mem):
        alloc[name]["cpu"] += cpu; alloc[name]["mem"] += mem

    desired = {}
    for dep in deployments:
        name = dep.get("name")
        if not name: continue
        cpu_req = _int(dep.get("cpu_req"), 100); mem_req = _int(dep.get("mem_req"), 64)
        placeable = [n["name"] for n in schedulable_fresh if _matches(n, dep)]   # affinity + taints
        for i in range(1, _int(dep.get("replicas"), 0) + 1):
            pname = f"{name}-{i}"
            prev = (existing.get(pname) or {}).get("node")
            chosen = ""
            if prev in keepable and _matches(by_name[prev], dep) and fits(prev, cpu_req, mem_req):
                chosen = prev
            else:
                candidates = [nm for nm in placeable if fits(nm, cpu_req, mem_req)]
                if candidates:
                    chosen = max(candidates, key=lambda nm: cap[nm][0] - alloc[nm]["cpu"])
            if chosen: place(chosen, cpu_req, mem_req)
            desired[pname] = {"deployment": name, "node": chosen,
                              "image": dep.get("image"), "command": dep.get("command") or "",
                              "cpu_req": cpu_req, "mem_req": mem_req}
    return desired, alloc

def _load_nodes(ns, now):
    out = []
    for r in _dicts(ns):
        hb = _int(r.get("last_heartbeat"), 0)
        out.append({"name": r["name"], "cpu_total": _int(r.get("cpu_total"), 1000),
                    "mem_total": _int(r.get("mem_total"), 512),
                    "fresh": now - hb < NODE_TTL,
                    "schedulable": _truthy(r.get("schedulable")),
                    "labels": parse_kv(r.get("labels")),
                    "taints": parse_taints(r.get("taints"))})
    return out

def _load_deployments(ws):
    out = []
    for d in _dicts(ws):
        d = dict(d)
        d["node_selector"] = parse_kv(d.get("node_selector"))
        d["tolerations"] = parse_tolerations(d.get("tolerations"))
        out.append(d)
    return out

def _write_pods(ps, desired, existing, rep):
    if ps.max_row > 1: ps.delete_rows(2, ps.max_row - 1)
    for pname, d in desired.items():
        live = rep.get(pname) or {}
        if not d["node"]:
            phase = "Unschedulable"
        else:
            phase = live.get("phase") or ("Running" if pname in rep else "Pending")
        cid = live.get("container_id") or (existing.get(pname) or {}).get("container_id") or ""
        ps.append([pname, d["deployment"], d["node"], phase, cid])

def _write_node_status(ns, nodes, alloc):
    ci = {h: TABS["Nodes"].index(h) for h in ("cpu_used", "status", "schedulable")}
    for row in ns.iter_rows(min_row=2):
        nm = row[0].value
        n = next((x for x in nodes if x["name"] == nm), None)
        if not n: continue
        row[ci["cpu_used"]].value = alloc.get(nm, {}).get("cpu", 0)
        if not n["fresh"]:            row[ci["status"]].value = "NotReady"
        elif not n["schedulable"]:    row[ci["status"]].value = "SchedulingDisabled"
        else:                         row[ci["status"]].value = "Ready"

def _reschedule(wb, exclude=frozenset()):
    """Recompute placement from current state and persist Pods + node status.
    Used by control actions (cordon/drain); no live kubelet report to fold in."""
    now = int(time.time()); ns = wb["Nodes"]
    nodes = _load_nodes(ns, now)
    existing = {p["name"]: p for p in _dicts(wb["Pods"])}
    desired, alloc = schedule(_load_deployments(wb["Deployments"]), nodes, existing, exclude)
    _write_pods(wb["Pods"], desired, existing, rep={})
    _write_node_status(ns, nodes, alloc)
    return desired

def heartbeat(node, ip, cpu_total, mem_total, reported):
    """Node reports in -> upsert it, run the scheduler over all Deployments, rewrite
    Pods, and return this node's desired pod set (same response shape as the Apps
    Script apiserver, so kubelet.sh is byte-for-byte identical)."""
    wb = _wb(); now = int(time.time())
    ns = wb["Nodes"]; nh = TABS["Nodes"]

    # 1) upsert the reporting node (preserve cpu_used + schedulable/cordon state).
    found = False
    for row in ns.iter_rows(min_row=2):
        if row[0].value == node:
            patch = {"ip": ip, "cpu_total": cpu_total, "mem_total": mem_total, "last_heartbeat": now}
            for i, h in enumerate(nh):
                if h in patch: row[i].value = patch[h]
            found = True; break
    if not found:
        ns.append([node, ip, cpu_total, 0, mem_total, "Ready", now, True])

    # 2) schedule over the current fleet.
    nodes = _load_nodes(ns, now)
    existing = {p["name"]: p for p in _dicts(wb["Pods"])}
    rep = {p.get("name"): p for p in (reported or [])}
    desired, alloc = schedule(_load_deployments(wb["Deployments"]), nodes, existing)

    # 3) persist Pods + node status/allocation.
    _write_pods(wb["Pods"], desired, existing, rep)
    _write_node_status(ns, nodes, alloc)
    wb.save(WORKBOOK)

    # 4) this node's marching orders: run what's assigned here, stop the rest.
    out = [{"name": p, "desired": "Running", "image": d["image"], "command": d["command"],
            "cpu_req": d["cpu_req"], "mem_req": d["mem_req"], "deployment": d["deployment"]}
           for p, d in desired.items() if d["node"] == node]
    for pname in rep:
        d = desired.get(pname)
        if not d or d["node"] != node:
            out.append({"name": pname, "desired": "Terminating"})
    return {"pods": out}

# ----------------------------------------------------------------------- HTTP

class H(BaseHTTPRequestHandler):
    def _json(self, obj, code=200):
        body = json.dumps(obj).encode()
        self.send_response(code); self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)
    def do_GET(self):
        q = parse_qs(urlparse(self.path).query)
        if q.get("token", [""])[0] != TOKEN: return self._json({"error": "unauthorized"}, 401)
        items = read_tab(q.get("kind", ["pods"])[0])
        if items is None: return self._json({"error": "unknown kind"}, 400)
        self._json({"items": items})
    def do_POST(self):
        n = int(self.headers.get("Content-Length", 0))
        raw = self.rfile.read(n) or b"{}"
        if SIGNING_KEY and not verify(SIGNING_KEY, self.headers.get("X-SNCF-Timestamp", ""),
                                      self.headers.get("X-SNCF-Signature", ""), raw,
                                      int(time.time()), SIGN_TTL):
            return self._json({"error": "bad signature"}, 401)
        try: body = json.loads(raw)
        except Exception: return self._json({"error": "bad json"}, 400)
        if body.get("token") != TOKEN: return self._json({"error": "unauthorized"}, 401)
        a = body.get("action")
        if a is None and body.get("node"):   # kubelet heartbeat
            return self._json(heartbeat(
                body.get("node"), body.get("ip", ""),
                _int(body.get("cpu_total"), 1000), _int(body.get("mem_total"), 512),
                body.get("pods", [])))
        if a == "apply":
            self._json({"applied": [upsert_deployment(d) for d in body.get("deployments", [])]})
        elif a == "scale":     self._json({"result": scale(body.get("name"), body.get("replicas"))})
        elif a == "delete":    self._json({"result": delete(body.get("name"))})
        elif a == "cordon":    self._json({"result": set_schedulable(body.get("name"), False)})
        elif a == "uncordon":  self._json({"result": set_schedulable(body.get("name"), True)})
        elif a == "drain":     self._json({"result": drain(body.get("name"))})
        elif a == "migrate":   self._json({"result": migrate(body.get("name"), body.get("node"))})
        elif a == "label":     self._json({"result": label_node(body.get("name"), body.get("spec", ""))})
        elif a == "taint":     self._json({"result": taint_node(body.get("name"), body.get("spec", ""))})
        else: self._json({"error": "unknown action"}, 400)
    def log_message(self, *a): pass

if __name__ == "__main__":
    _wb()
    print(f"[apiserver] serving {WORKBOOK} on http://0.0.0.0:{PORT} (kinds: {', '.join(TABS)})"
          + (" · HMAC signing REQUIRED on POST" if SIGNING_KEY else ""))
    ThreadingHTTPServer(("0.0.0.0", PORT), H).serve_forever()
