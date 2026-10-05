"""Published ports (the NodePort analog): parsing, host-port-aware scheduling,
endpoint reporting and the apiserver round trip. No Docker required."""
import os, sys, importlib
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pytest
from apiserver import (parse_ports, parse_docker_ports, format_ports, port_range,
                       schedule, pod_rows, node_orders)


def node(name, cpu=4000, mem=8192, fresh=True, schedulable=True):
    return {"name": name, "cpu_total": cpu, "mem_total": mem, "fresh": fresh,
            "schedulable": schedulable, "labels": {}, "taints": []}

def dep(name, replicas=1, ports="", cpu=100, mem=64):
    return {"name": name, "image": "nginx", "replicas": replicas, "cpu_req": cpu,
            "mem_req": mem, "command": "", "node_selector": {}, "tolerations": set(),
            "ports": ports}


# ---------------------------------------------------------------- parsing

def test_parse_ports_forms():
    assert parse_ports("80") == [(None, 80, "tcp")]
    assert parse_ports("8080:80") == [(8080, 80, "tcp")]
    assert parse_ports("8080:80, 53/udp") == [(8080, 80, "tcp"), (None, 53, "udp")]
    assert parse_ports(["8443:443/TCP"]) == [(8443, 443, "tcp")]
    assert parse_ports("") == [] and parse_ports(None) == []

@pytest.mark.parametrize("bad", ["http", "80:abc", "0", "70000", "8080:80/sctp", "1:2:3"])
def test_parse_ports_rejects_garbage(bad):
    with pytest.raises(ValueError):
        parse_ports(bad)

def test_format_and_range():
    assert format_ports([(30000, 80, "tcp"), (53, 53, "udp")]) == "30000:80/tcp,53:53/udp"
    assert port_range("31000-31010") == (31000, 31010)

def test_parse_docker_ports_dedups_ipv4_ipv6_and_ignores_exposed():
    s = "0.0.0.0:30000->80/tcp, [::]:30000->80/tcp, 443/tcp, 127.0.0.1:5353->53/udp"
    assert parse_docker_ports(s) == {(30000, 80, "tcp"), (5353, 53, "udp")}
    assert parse_docker_ports("") == set()


# ------------------------------------------------------------- scheduling

def test_fixed_host_port_spreads_replicas_across_nodes():
    desired, _ = schedule([dep("gw", replicas=2, ports="8080:80")], [node("a"), node("b")], {})
    assert {d["node"] for d in desired.values()} == {"a", "b"}
    assert all(d["ports"] == "8080:80/tcp" for d in desired.values())

def test_fixed_host_port_conflict_makes_extra_replica_unschedulable():
    desired, alloc = schedule([dep("gw", replicas=2, ports="8080:80")], [node("a")], {})
    assert desired["gw-1"]["node"] == "a"
    assert desired["gw-2"]["node"] == ""                 # port 8080 already bound on a
    assert alloc["a"]["cpu"] == 100                      # the loser consumes nothing

def test_conflict_across_deployments_and_protocols():
    desired, _ = schedule([dep("x", ports="8080:80"), dep("y", ports="8080:8080"),
                           dep("dns", ports="8080:53/udp")], [node("a")], {})
    assert desired["x-1"]["node"] == "a"
    assert desired["y-1"]["node"] == ""                  # same tcp host port
    assert desired["dns-1"]["node"] == "a"               # udp is a different socket

def test_auto_ports_are_distinct_on_one_node():
    desired, _ = schedule([dep("web", replicas=3, ports="80")], [node("a")], {}, nodeports="30000-30010")
    assert [d["ports"] for d in desired.values()] == ["30000:80/tcp", "30001:80/tcp", "30002:80/tcp"]

def test_auto_ports_never_take_a_requested_fixed_port():
    desired, _ = schedule([dep("web", ports="80"), dep("pinned", ports="30000:80")],
                          [node("a")], {}, nodeports="30000-30010")
    assert desired["web-1"]["ports"] == "30001:80/tcp"
    assert desired["pinned-1"]["ports"] == "30000:80/tcp"

def test_auto_port_range_exhaustion_is_unschedulable():
    desired, _ = schedule([dep("web", replicas=3, ports="80")], [node("a")], {}, nodeports="30000-30001")
    assert [d["node"] for d in desired.values()] == ["a", "a", ""]

def test_auto_port_is_sticky_across_reschedules():
    existing = {"web-1": {"node": "a", "ports": "30007:80/tcp"}}
    desired, _ = schedule([dep("web", replicas=2, ports="80")], [node("a")], existing,
                          nodeports="30000-30010")
    assert desired["web-1"]["ports"] == "30007:80/tcp"   # kept its port
    assert desired["web-2"]["ports"] == "30000:80/tcp"

def test_running_pod_keeps_its_port_against_an_earlier_newcomer():
    # 'new' comes first in row order, but 'old' already holds 8080 on a: two-phase
    # placement keeps the running pod and sends the newcomer elsewhere.
    existing = {"old-1": {"node": "a", "ports": "8080:80/tcp"}}
    desired, _ = schedule([dep("new", ports="8080:80"), dep("old", ports="8080:80")],
                          [node("a"), node("b")], existing)
    assert desired["old-1"]["node"] == "a"
    assert desired["new-1"]["node"] == "b"

def test_spec_hash_tracks_ports():
    d1, _ = schedule([dep("web", ports="8080:80")], [node("a")], {})
    d2, _ = schedule([dep("web", ports="9090:80")], [node("a")], {})
    d3, _ = schedule([dep("web", ports="8080:80")], [node("a")], {})
    assert d1["web-1"]["spec_hash"] != d2["web-1"]["spec_hash"]
    assert d1["web-1"]["spec_hash"] == d3["web-1"]["spec_hash"]

def test_garbage_ports_cell_does_not_break_scheduling():
    desired, _ = schedule([dep("web", ports="banana")], [node("a")], {})
    assert desired["web-1"]["node"] == "a" and desired["web-1"]["ports"] == ""


# ------------------------------------------------------- status & orders

def test_pod_rows_report_endpoints_from_kubelet():
    desired, _ = schedule([dep("web", ports="8080:80")], [node("a")], {})
    rows = pod_rows(desired, {}, {"web-1": {"phase": "Running", "container_id": "abc",
                                            "ports": "0.0.0.0:8080->80/tcp, [::]:8080->80/tcp"}},
                    {"a": "10.0.0.5"}, reporter="a")
    assert rows[0]["endpoints"] == "10.0.0.5:8080->80/tcp"
    assert rows[0]["ports"] == "8080:80/tcp" and rows[0]["phase"] == "Running"

def test_pod_rows_keep_other_nodes_status():
    desired, _ = schedule([dep("web", replicas=2, ports="8080:80")], [node("a"), node("b")], {})
    on_b = next(p for p, d in desired.items() if d["node"] == "b")
    existing = {on_b: {"node": "b", "phase": "Running", "endpoints": "10.0.0.6:8080->80/tcp",
                       "container_id": "xyz"}}
    rows = {r["name"]: r for r in pod_rows(desired, existing, {}, {}, reporter="a")}
    assert rows[on_b]["phase"] == "Running"                  # a can't vouch for b's pod
    assert rows[on_b]["endpoints"] == "10.0.0.6:8080->80/tcp"

def test_node_orders_carry_ports_and_spec_hash():
    desired, _ = schedule([dep("web", ports="8080:80")], [node("a")], {})
    out = node_orders(desired, {}, "a")
    assert out[0]["ports"] == "8080:80/tcp" and len(out[0]["spec_hash"]) == 12


# ------------------------------------------------------ apiserver round trip

@pytest.fixture
def api(tmp_path):
    os.environ["WORKBOOK"] = str(tmp_path / "cluster.xlsx")
    os.environ["TOKEN"] = "t"
    import apiserver
    importlib.reload(apiserver)
    return apiserver

def test_apply_rejects_bad_ports(api):
    assert api.upsert_deployment({"name": "web", "image": "nginx", "replicas": 1,
                                  "ports": "eighty"}).startswith("invalid")
    assert api.read_tab("deployments") == []

def test_apply_accepts_port_list(api):
    assert api.upsert_deployment({"name": "web", "image": "nginx", "replicas": 1,
                                  "ports": ["8080:80", "53/udp"]}) == "created"
    assert api.read_tab("deployments")[0]["ports"] == "8080:80,53/udp"

def test_heartbeat_publishes_and_reports(api):
    api.upsert_deployment({"name": "web", "image": "nginx", "replicas": 2, "cpu_req": 100,
                           "mem_req": 64, "ports": "8080:80"})
    resp = api.heartbeat("a", "10.0.0.5", 4000, 8192, [])
    run = [p for p in resp["pods"] if p["desired"] == "Running"]
    assert [p["ports"] for p in run] == ["8080:80/tcp"]      # second replica can't bind 8080
    pods = {p["name"]: p for p in api.read_tab("pods")}
    assert pods["web-2"]["phase"] == "Unschedulable"
    api.heartbeat("a", "10.0.0.5", 4000, 8192,
                  [{"name": "web-1", "container_id": "c1", "ports": "0.0.0.0:8080->80/tcp"}])
    pods = {p["name"]: p for p in api.read_tab("pods")}
    assert pods["web-1"]["endpoints"] == "10.0.0.5:8080->80/tcp"
    # a second node joins: the pending replica gets it, with the same host port
    api.heartbeat("b", "10.0.0.6", 4000, 8192, [])
    pods = {p["name"]: p for p in api.read_tab("pods")}
    assert pods["web-2"]["node"] == "b" and pods["web-2"]["ports"] == "8080:80/tcp"
    assert pods["web-1"]["endpoints"] == "10.0.0.5:8080->80/tcp"   # untouched by b's report

def test_schema_upgrade_adds_ports_columns(tmp_path):
    import openpyxl
    wbpath = tmp_path / "old.xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    wb.create_sheet("Deployments").append(["name", "image", "replicas", "cpu_req", "mem_req", "command"])
    wb["Deployments"].append(["web", "nginx", 1, 100, 64, ""])
    wb.create_sheet("Pods").append(["name", "deployment", "node", "phase", "container_id"])
    wb.save(wbpath)
    os.environ["WORKBOOK"] = str(wbpath); os.environ["TOKEN"] = "t"
    import apiserver; importlib.reload(apiserver)
    resp = apiserver.heartbeat("a", "1", 4000, 8192, [])
    assert resp["pods"][0]["ports"] == ""
    assert "endpoints" in apiserver.read_tab("pods")[0]
